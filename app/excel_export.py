from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from pathlib import Path

HEADERS=["SL NO","AREA","LINE NO.","DRAWING NO.","REV.","SHEET","ITEM","SIZE (IN)","SCHEDULE/CLASS","MOC","GRADE","DESCRIPTION","QTY","UOM","REMARKS"]

def style(ws, widths=None):
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:
        c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="0F172A"); c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    for row in ws.iter_rows():
        for c in row[0:]: c.alignment=Alignment(vertical="top",wrap_text=True)
    if widths:
        for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w

def write_rows(ws, headers, rows):
    ws.append(headers)
    for r in rows: ws.append([r.get(h,"") for h in headers])
    style(ws,[8,18,26,28,8,12,20,12,18,14,24,65,10,10,32][:len(headers)])

def export_bom(result,path:Path,template_path=None):
    wb=Workbook(); ws=wb.active; ws.title="FINAL BOM"
    write_rows(ws,HEADERS,result["rows"])
    fab=[r for r in result["rows"] if "FABRICATION" in str(r.get("REMARKS","" )).upper()]
    ere=[r for r in result["rows"] if "ERECTION" in str(r.get("REMARKS","" )).upper()]
    wf=wb.create_sheet("FABRICATION BOM"); write_rows(wf,HEADERS,fab)
    we=wb.create_sheet("ERECTION BOM"); write_rows(we,HEADERS,ere)
    cut_headers=["AREA","LINE NO.","DRAWING NO.","REV.","SHEET","PIECE NO","CUT LENGTH","N.S.","REMARKS"]
    wc=wb.create_sheet("CUT PIPE"); write_rows(wc,cut_headers,result.get("cut_pipe",[]))
    wr=wb.create_sheet("REVIEW REQUIRED");
    wr.append(["ROW","REASON","DRAWING NO.","SHEET","ITEM","SIZE","QTY","DESCRIPTION"])
    for x in result["review"]:
        d=x.get("data",{}); wr.append([x.get("row"),x.get("reason"),d.get("DRAWING NO."),d.get("SHEET"),d.get("ITEM"),d.get("SIZE (IN)"),d.get("QTY"),d.get("DESCRIPTION")])
    style(wr,[8,45,28,12,20,12,10,70])
    wsr=wb.create_sheet("VALIDATION REPORT");
    wsr.append(["METRIC","VALUE"]); wsr.append(["Total BOM rows",len(result["rows"])]); wsr.append(["Review flags",len(result["review"])]); wsr.append(["Cut pipe rows",len(result.get("cut_pipe",[]))]); wsr.append(["Pages",result["pages"]]); wsr.append(["Document type",result["document_type"]]);
    for w in result.get("warnings",[]): wsr.append(["Warning",w])
    style(wsr,[28,80])
    src=wb.create_sheet("SOURCE OCR")
    src.append(["PAGE","OCR TEXT"])
    for i,t in enumerate(result["page_texts"],1): src.append([i,t])
    style(src,[10,120])
    wb.save(path)
